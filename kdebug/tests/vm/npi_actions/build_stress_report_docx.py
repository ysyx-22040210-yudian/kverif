#!/usr/bin/env python3
"""Build the Chinese KDebug NPI VM stress-test Word report from evidence files."""

from __future__ import print_function

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PRESET = {
    "name": "compact_reference_guide",
    "page_width": 12240,
    "page_height": 15840,
    "margin": 1440,
    "header_distance": 708,
    "footer_distance": 708,
    "content_width": 9360,
    "table_indent": 120,
    "cell_top": 80,
    "cell_bottom": 80,
    "cell_start": 120,
    "cell_end": 120,
    "body_after": 120,
    "body_line": 300,
    "h1_before": 360,
    "h1_after": 200,
    "h2_before": 280,
    "h2_after": 140,
    "h3_before": 200,
    "h3_after": 100,
}

COLORS = {
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "ink": "0B2545",
    "muted": "5F6B76",
    "grid": "B8C2CC",
    "header_fill": "E8EEF5",
    "light_fill": "F4F6F9",
    "pass_fill": "E7F3EA",
    "pass_text": "205D35",
    "warn_fill": "FFF3CD",
    "warn_text": "7A5A00",
    "risk": "9B1C1C",
    "white": "FFFFFF",
}

ACTION_ORDER = [
    "npi.capabilities",
    "language.resolve", "language.iterate", "language.relate", "language.value",
    "module.find_instances", "module.inspect", "module.objects",
    "netlist.resolve", "netlist.iterate",
    "text.line", "text.words", "text.replace_line",
    "dm.add_net", "dm.clone_module", "vcs.summary",
    "power.resolve", "power.list",
    "crdb.resolve", "crdb.correlates",
    "transaction.writer.create", "fsdb.writer.create_scope",
]

MODULE_ORDER = [
    "continuous_assignments", "functions", "generate_scopes", "instances",
    "instances_in_generate", "io", "language_interfaces", "nets", "parameters",
    "ports", "primitives", "always_processes", "initial_processes", "tasks", "variables",
]


def set_run_font(run, size=None, bold=None, color=None, ascii_font="Calibri",
                 east_asia_font="Microsoft YaHei", italic=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=240):
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def set_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def set_paragraph_shading(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "10")
        left.set(qn("w:color"), left_border)
        borders.append(left)


def set_paragraph_bottom_border(paragraph, color="2E74B5", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_page_field(paragraph, field_name):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=8.5, color=COLORS["muted"])


def set_cell_text(cell, text, size=8.2, bold=False, color="0B2545", align="left",
                  ascii_font="Calibri"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    set_paragraph_spacing(paragraph, before=0, after=0, line=250)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color, ascii_font=ascii_font)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths):
    assert sum(widths) == PRESET["content_width"]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    for tag, value, attr in (
        ("w:tblW", PRESET["content_width"], "w:w"),
        ("w:tblInd", PRESET["table_indent"], "w:w"),
    ):
        element = tbl_pr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_pr.append(element)
        element.set(qn(attr), str(value))
        element.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in (("top", PRESET["cell_top"]), ("bottom", PRESET["cell_bottom"]),
                        ("start", PRESET["cell_start"]), ("end", PRESET["cell_end"])):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn("w:" + edge))
        if border is None:
            border = OxmlElement("w:" + edge)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), COLORS["grid"])

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style="Heading %d" % level)
    paragraph.add_run(text)
    set_keep(paragraph, keep_next=True)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True, color=COLORS["ink"])
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=COLORS["ink"])
    else:
        run = paragraph.add_run(text)
        set_run_font(run, color=COLORS["ink"])
    return paragraph


def add_callout(doc, title, body, kind="info"):
    paragraph = doc.add_paragraph()
    fill = COLORS["warn_fill"] if kind == "warning" else COLORS["light_fill"]
    border = COLORS["warn_text"] if kind == "warning" else COLORS["blue"]
    set_paragraph_shading(paragraph, fill, border)
    set_paragraph_spacing(paragraph, before=80, after=120, line=280)
    paragraph.paragraph_format.left_indent = Twips(180)
    paragraph.paragraph_format.right_indent = Twips(120)
    heading = paragraph.add_run(title + "  ")
    set_run_font(heading, size=10.5, bold=True, color=border)
    run = paragraph.add_run(body)
    set_run_font(run, size=10.5, color=COLORS["ink"])
    return paragraph


def add_code_block(doc, text):
    paragraph = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(paragraph, "F2F4F7", "7A8793")
    paragraph.paragraph_format.left_indent = Twips(180)
    paragraph.paragraph_format.right_indent = Twips(120)
    run = paragraph.add_run(text)
    set_run_font(run, size=8.2, ascii_font="Consolas", east_asia_font="Microsoft YaHei",
                 color="263238")
    return paragraph


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.rows[0]._element.getparent().remove(table.rows[0]._element)
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, size=9, bold=True, color=COLORS["dark_blue"])
        shade_cell(cells[0], COLORS["header_fill"])
        set_cell_text(cells[1], value, size=9, color=COLORS["ink"])
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, [2700, 6660])
    return table


def add_action_table(doc, rows):
    widths = [2360, 540, 540, 540, 620, 500, 1860, 2400]
    table = doc.add_table(rows=1, cols=len(widths))
    headers = ["Action", "项", "次数", "通过", "阻塞", "失败", "耗时 avg / p95 / max (s)", "状态"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_text(cell, value, size=8.2, bold=True, color=COLORS["ink"], align="center")
        shade_cell(cell, COLORS["header_fill"])
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        values = [
            row["action"], row["case_count"], row["attempts"], row["passed"],
            row["license_blocked"], row["failed"],
            "%.3f / %.3f / %.3f" % (float(row["elapsed_avg_sec"]),
                                     float(row["elapsed_p95_sec"]),
                                     float(row["elapsed_max_sec"])),
            row["status"],
        ]
        for index, (cell, value) in enumerate(zip(cells, values)):
            align = "left" if index == 0 else "center"
            status_color = COLORS["ink"]
            bold = index in (0, 7)
            if index == 7:
                if row["status"] == "PASS":
                    status_color = COLORS["pass_text"]
                    shade_cell(cell, COLORS["pass_fill"])
                else:
                    status_color = COLORS["warn_text"]
                    shade_cell(cell, COLORS["warn_fill"])
            set_cell_text(cell, value, size=7.9, bold=bold, color=status_color, align=align,
                          ascii_font="Consolas" if index == 0 else "Calibri")
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths)
    return table


def add_module_table(doc, rows):
    widths = [3000, 650, 650, 600, 2060, 2400]
    table = doc.add_table(rows=1, cols=len(widths))
    headers = ["module.objects kind", "Runs", "PASS", "Fail", "耗时 avg / p95 / max (s)", "结果"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_text(cell, value, size=8.2, bold=True, color=COLORS["ink"], align="center")
        shade_cell(cell, COLORS["header_fill"])
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row in rows:
        result = row["status"]
        if row["kind"] == "language_interfaces":
            result = "PASS（纯 SV 预期 0 项）"
        cells = table.add_row().cells
        values = [
            row["kind"], row["attempts"], row["passed"], row["failed"],
            "%.3f / %.3f / %.3f" % (float(row["elapsed_avg_sec"]),
                                     float(row["elapsed_p95_sec"]),
                                     float(row["elapsed_max_sec"])),
            result,
        ]
        for index, (cell, value) in enumerate(zip(cells, values)):
            set_cell_text(cell, value, size=7.9, bold=index in (0, 5),
                          color=COLORS["pass_text"] if index == 5 else COLORS["ink"],
                          align="left" if index == 0 else "center",
                          ascii_font="Consolas" if index == 0 else "Calibri")
            if index == 5:
                shade_cell(cell, COLORS["pass_fill"])
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths)
    return table


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_text(cell, value, size=8.7, bold=True, color=COLORS["ink"], align="center")
        shade_cell(cell, COLORS["header_fill"])
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            set_cell_text(cell, value, size=8.5, color=COLORS["ink"],
                          align="left" if index == len(cells) - 1 else "center")
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal_r_pr = normal._element.get_or_add_rPr()
    normal_r_pr.rFonts.set(qn("w:ascii"), "Calibri")
    normal_r_pr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal_r_pr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal_p_pr = normal._element.get_or_add_pPr()
    normal_spacing = OxmlElement("w:spacing")
    normal_spacing.set(qn("w:before"), "0")
    normal_spacing.set(qn("w:after"), str(PRESET["body_after"]))
    normal_spacing.set(qn("w:line"), str(PRESET["body_line"]))
    normal_spacing.set(qn("w:lineRule"), "auto")
    normal_p_pr.append(normal_spacing)

    heading_tokens = {
        1: (16, COLORS["blue"], PRESET["h1_before"], PRESET["h1_after"]),
        2: (13, COLORS["blue"], PRESET["h2_before"], PRESET["h2_after"]),
        3: (12, COLORS["dark_blue"], PRESET["h3_before"], PRESET["h3_after"]),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles["Heading %d" % level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        p_pr = style._element.get_or_add_pPr()
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), str(before))
        spacing.set(qn("w:after"), str(after))
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        p_pr.append(OxmlElement("w:keepNext"))
        p_pr.append(OxmlElement("w:keepLines"))

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.2)
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    p_pr = code._element.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "80")
    spacing.set(qn("w:after"), "120")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)


def configure_page(section):
    section.page_width = Twips(PRESET["page_width"])
    section.page_height = Twips(PRESET["page_height"])
    section.top_margin = Twips(PRESET["margin"])
    section.bottom_margin = Twips(PRESET["margin"])
    section.left_margin = Twips(PRESET["margin"])
    section.right_margin = Twips(PRESET["margin"])
    section.header_distance = Twips(PRESET["header_distance"])
    section.footer_distance = Twips(PRESET["footer_distance"])


def configure_header_footer(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, before=0, after=0, line=240)
    left = paragraph.add_run("KVERIF  |  KDEBUG NPI VM 压测")
    set_run_font(left, size=8.5, bold=True, color=COLORS["muted"])
    right = paragraph.add_run("    2026-07-29")
    set_run_font(right, size=8.5, color=COLORS["muted"])

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(paragraph, before=0, after=0, line=240)
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=COLORS["muted"])
    add_page_field(paragraph, "PAGE")
    run = paragraph.add_run(" 页 / 共 ")
    set_run_font(run, size=8.5, color=COLORS["muted"])
    add_page_field(paragraph, "NUMPAGES")
    run = paragraph.add_run(" 页")
    set_run_font(run, size=8.5, color=COLORS["muted"])


def load_data(evidence):
    with open(evidence / "stress-summary.json", encoding="utf-8") as stream:
        summary = json.load(stream)
    with open(evidence / "stress-action-results.csv", encoding="utf-8", newline="") as stream:
        action_rows = list(csv.DictReader(stream))
    with open(evidence / "stress-results.csv", encoding="utf-8", newline="") as stream:
        case_rows = list(csv.DictReader(stream))
    action_by_name = {row["action"]: row for row in action_rows}
    kind_by_name = {row["kind"]: row for row in case_rows if row["kind"]}
    action_rows = [action_by_name[name] for name in ACTION_ORDER]
    module_rows = [kind_by_name[name] for name in MODULE_ORDER]
    assert summary["scope"]["unique_actions"] == 22
    assert summary["scope"]["independent_cases"] == 36
    assert summary["totals"]["attempts"] == 360
    assert summary["totals"]["failed"] == 0
    assert len(action_rows) == 22 and len(module_rows) == 15
    return summary, action_rows, module_rows


def build_document(evidence, output):
    summary, action_rows, module_rows = load_data(evidence)
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)

    props = doc.core_properties
    props.title = "KDebug Tcl NPI Action VM 压测报告"
    props.subject = "Verdi O-2018.09-SP2 public CLI stress test"
    props.author = "KVerif Project"
    props.keywords = "KVerif,KDebug,NPI,Verdi,VCS,VM,stress test"
    props.comments = "Generated from archived machine-readable evidence."

    kicker = doc.add_paragraph()
    set_paragraph_spacing(kicker, before=120, after=100, line=240)
    run = kicker.add_run("TECHNICAL VALIDATION REPORT")
    set_run_font(run, size=9.5, bold=True, color=COLORS["blue"])

    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=0, after=80, line=340)
    run = title.add_run("KDebug Tcl NPI Action VM 压测报告")
    set_run_font(run, size=23, bold=True, color=COLORS["ink"])

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, before=0, after=180, line=280)
    run = subtitle.add_run("Verdi O-2018.09-SP2  |  host 普通用户  |  10 次/项  |  2 路并发")
    set_run_font(run, size=11.5, color=COLORS["muted"])

    metadata = doc.add_paragraph()
    set_paragraph_spacing(metadata, before=0, after=180, line=260)
    set_paragraph_bottom_border(metadata)
    for label, value in (
        ("测试日期", "2026-07-29"),
        ("VM", "host@192.168.31.116"),
        ("结果目录", "/home/host/kverif_npi_action_stress_20260729"),
    ):
        label_run = metadata.add_run(label + "：")
        set_run_font(label_run, size=9.5, bold=True, color=COLORS["dark_blue"])
        value_run = metadata.add_run(value + "\n")
        set_run_font(value_run, size=9.5, color=COLORS["ink"],
                     ascii_font="Consolas" if "/" in value or "@" in value else "Calibri")

    add_callout(
        doc,
        "核心结论",
        "360 次受测 action 调用得到 340 PASS、20 次 LICENSE_BLOCKED、0 次非预期失败；"
        "另有 10 次 public scope.list 重开校验。20/22 个 action 完成功能压测，两个 Power action 因许可证阻塞不能计为 PASS。",
    )

    add_heading(doc, "1. 执行结论", 1)
    add_key_value_table(doc, [
        ("覆盖范围", "22 个新增 action，拆为 36 个独立压测项"),
        ("执行规模", "每项 10 次、并发度 2；360 次受测调用，370 次 public CLI 总调用"),
        ("通过情况", "340 PASS；20 个 action 完成真实功能压测"),
        ("许可证阻塞", "20 次；power.resolve 和 power.list 各 10 次"),
        ("非预期失败", "0"),
        ("总墙钟时间", "%.6f 秒" % summary["run"]["wall_elapsed_sec"]),
        ("调用边界", "只调用 tools/kdebug；直接 Tcl/NPI 调用为 0"),
        ("退出状态", "accepted=true；含许可证阻塞，不等同于 22/22 功能 PASS"),
    ])

    doc.add_page_break()
    add_heading(doc, "2. 测试环境与口径", 1)
    env = summary["environment"]
    add_key_value_table(doc, [
        ("执行用户", env.get("user", "host")),
        ("VM 主机", env.get("hostname", "") + "（192.168.31.116）"),
        ("Verdi", env.get("verdi_home", "")),
        ("VCS", env.get("vcs_home", "") + "；vcs.summary 报告 O-2018.09-1"),
        ("Python", env.get("python", "")),
        ("KDebug 入口", env.get("kdebug", "")),
        ("Wrapper SHA-256", env.get("kdebug_sha256", "")),
        ("Harness timeout", "未设置；单次 action 不由压测 harness 截断"),
    ])

    add_body(
        doc,
        "测试使用 tiny VCS -kdb/-Xdump_vcsdb 设计、Verdi 自带 RTL+UPF Power demo、真实 RTL/GATE CRDB，"
        "以及 transaction/signal FSDB writer。修改类和 writer action 每轮使用唯一输出路径；两个并发进程使用独立工作目录。",
    )

    add_heading(doc, "3. 实际执行命令", 1)
    command = (evidence / "stress-command.txt").read_text(encoding="utf-8").strip()
    formatted_command = command.replace(" KDEBUG_STRESS_PARALLEL", " \\\n+KDEBUG_STRESS_PARALLEL")
    formatted_command = formatted_command.replace(" KVERIF_HOME", " \\\n+KVERIF_HOME")
    formatted_command = formatted_command.replace(" KDEBUG_BIN", " \\\n+KDEBUG_BIN")
    formatted_command = formatted_command.replace(" bash ", " \\\n+bash ")
    formatted_command = formatted_command.replace(" /home/host/kverif_npi_action_stress_20260729", " \\\n+  /home/host/kverif_npi_action_stress_20260729")
    formatted_command = formatted_command.replace(chr(10) + "+", chr(10))
    add_code_block(doc, formatted_command)
    add_body(doc, "输出目录必须事先不存在；复测时应改用新的时间戳目录。", bold_prefix="输出目录")

    doc.add_page_break()
    add_heading(doc, "4. 22 个 action 压测结果", 1)
    add_body(doc, "耗时列依次为平均值、p95 和最大值，单位为秒。module.objects 汇总 15 个 kind，共 150 次。")
    add_action_table(doc, action_rows)

    add_callout(
        doc,
        "Power 结果口径",
        "power.resolve 与 power.list 均真实启动 Verdi 并加载 RTL+UPF，但 10/10 次在 PowerAwareAnalysis feature checkout 阶段返回 LICENSE_UNAVAILABLE。"
        "这证明错误处理路径稳定，不证明 Power 查询功能通过。",
        kind="warning",
    )

    doc.add_page_break()
    add_heading(doc, "5. module.objects 15 个 kind 明细", 1)
    add_body(doc, "每个 kind 独立执行 10 次。language_interfaces 在纯 SystemVerilog fixture 中稳定返回 0 项，该结果是预期值。")
    add_module_table(doc, module_rows)

    add_heading(doc, "6. 每轮语义断言", 1)
    add_matrix_table(doc, ["域", "机器断言"], [
        ("Language / Module", "WIDTH=12、BIAS=1、RESULT_WIDTH=12；实例定义和完整层次名一致。"),
        ("Port / IO", "lhs/rhs 为 input、result 为 output；3 个端口的 high/low connection 均非空。"),
        ("Module getters", "direct/generate 实例、net、function、task、primitive、always/initial 等 fixture 对象存在。"),
        ("Netlist", "npi_fixture_top.result[11:0] 宽度为 12；iterator 返回 7 个 fixture net。"),
        ("Text", "第 27 行与 token 正确；替换副本第 27 行为 sum = lhs - rhs;。"),
        ("Design Manipulation", "新增 [7:0] net 和 clone module 的导出目录包含非空 RTL 文件。"),
        ("VCS", "编译错误数为 0、模块数为 3、版本属于 O-2018.09。"),
        ("CRDB", "RTL state 可解析，并至少关联一个 GATE 对象。"),
        ("Transaction writer", "FSDB 非空、2 个 transaction、1 个 relation、结束时间为 45。"),
        ("FSDB writer", "FSDB 非空、3 个 scope/1 次 up；public scope.list 重开得到 top.u_a/top.u_b。"),
    ], [2300, 7060])

    add_heading(doc, "7. 性能观察", 1)
    add_body(
        doc,
        "所有通过项均在 6.729 秒内完成。最慢的 action 级最大值是 netlist.resolve 6.728 秒；"
        "module.objects 中最慢的 kind 最大值是 initial_processes 6.259 秒。没有 action timeout、进程崩溃、JSON 解析失败、"
        "输出覆盖冲突或结果漂移。",
    )

    add_heading(doc, "8. 验证与证据", 1)
    add_matrix_table(doc, ["证据文件", "内容"], [
        ("stress-summary.json", "环境、范围、总数、22 action 和 36 case 的结构化汇总。"),
        ("stress-action-results.csv", "22 个 action 的次数、状态及 min/avg/p50/p95/max。"),
        ("stress-results.csv", "36 个独立 case，含 15 个 module.objects kind。"),
        ("stress-attempts.jsonl", "360 次调用的命令、退出码、状态、耗时、响应哈希和日志路径。"),
        ("stress-command.txt", "实际执行命令。"),
        ("stress-implementation-sha256.txt", "Wrapper、binary、Tcl engine、harness 和 fixture 哈希。"),
    ], [3300, 6060])
    add_body(doc, "本地归档目录：E:\\xverif\\kdebug\\tests\\vm\\npi_actions\\evidence")
    add_body(doc, "VM 原始目录：/home/host/kverif_npi_action_stress_20260729")

    doc.add_page_break()
    add_heading(doc, "9. 其他回归结果", 1)
    add_key_value_table(doc, [
        ("Action contract", "107/107 通过"),
        ("Schema", "224/224 通过"),
        ("Examples", "219/219 通过"),
        ("CLI-only 工作流", "9/9 通过；覆盖 Bash、csh、Perl、Python、仓库外目录和空格路径"),
        ("Wrapper 自动发现", "清空 VERDI_HOME/NPIL1_PATH/LD_LIBRARY_PATH 后，module.objects ports=3"),
        ("残留进程", "0；本轮结束后无 KDebug/Verdi/VCS 测试进程"),
    ])

    add_heading(doc, "10. 最终判定与后续准入", 1)
    add_callout(
        doc,
        "当前判定",
        "20/22 个新增 action 已完成 VM 功能压测；两个 Power action 已完成重复调用和许可证错误路径压测，但未完成 Power 功能压测。",
        kind="warning",
    )
    add_body(
        doc,
        "只有 VM 获得 PowerAwareAnalysis license，并将 power.resolve 与 power.list 各重跑至 10/10 PASS 后，"
        "才能将总体结论升级为 22/22 action 完成功能压测。该复测不需要修改 KDebug 或 harness。",
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def parse_args():
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence", type=Path, default=base / "evidence")
    parser.add_argument("--output", type=Path,
                        default=base / "evidence" / "KDebug_NPI_VM_Stress_Test_Report_20260729.docx")
    return parser.parse_args()


def main():
    args = parse_args()
    output = build_document(args.evidence.resolve(), args.output.resolve())
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
